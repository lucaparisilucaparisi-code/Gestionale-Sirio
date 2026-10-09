"""Test dell'esportazione Excel (sirio/excel_export.py).

Dati di fantasia (tests/genera_esempio_excel.py). Il controllo del ricalcolo
con LibreOffice viene saltato se ``soffice`` non è installato; il test sul
foglio reale legge la verità attesa da SIRIO_SAMPLE_TRUTH e viene saltato se
la variabile non è impostata.
"""

from __future__ import annotations

import itertools
import json
import os
import re
import shutil
import subprocess
from datetime import datetime, time
from pathlib import Path

import openpyxl
import pytest
import xlsxwriter

from sirio import excel_export as ex
from sirio import validation as val
from sirio.excel_export import ExportOptions, default_filename, export_workbook
from sirio.models import DayRow, Document, Header, empty_rows
from tests.genera_esempio_excel import documenti_esempio

FISSI = ["Riepilogo", "Dettaglio giornaliero", "Anomalie", "Totali per operatore", "Legenda e note"]
RE_NOME_VIETATO = re.compile(r"[\[\]:*?/\\]")
RE_NUMERO_COME_TESTO = re.compile(r"^\s*-?\d+(?:[.,]\d+)?\s*$|^\s*\d{1,2}[:.]\d{2}\s*$")


# --- Supporto -------------------------------------------------------------------


@pytest.fixture(scope="module")
def esempio(tmp_path_factory) -> tuple[Path, list[Document]]:
    docs = documenti_esempio()
    path = export_workbook(docs, tmp_path_factory.mktemp("xlsx") / "esempio.xlsx")
    return path, docs


@pytest.fixture(scope="module")
def wb_formule(esempio) -> openpyxl.Workbook:
    return openpyxl.load_workbook(esempio[0], data_only=False)


@pytest.fixture(scope="module")
def wb_valori(esempio) -> openpyxl.Workbook:
    return openpyxl.load_workbook(esempio[0], data_only=True)


def validati(docs: list[Document]) -> dict[str, Document]:
    """Copie validate dei documenti inclusi, per id."""
    out = {}
    for d in docs:
        if d.status == "completato" and d.is_foglio_firma:
            c = d.model_copy(deep=True)
            val.validate_document(c)
            out[d.id] = c
    return out


def schede(wb: openpyxl.Workbook) -> list[str]:
    return [n for n in wb.sheetnames if n not in FISSI]


def trova_riga(ws, colonna: int, valore) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=colonna).value == valore:
            return r
    raise AssertionError(f"{valore!r} non trovato nella colonna {colonna} di {ws.title}")


def cella_scheda(ws, giorno: int, col0: int) -> openpyxl.cell.Cell:
    """Cella della tabella giornaliera di una scheda (col0 = indice 0-based come in _Esportatore)."""
    return ws.cell(row=ex._Esportatore.S_R_D1 + giorno, column=col0 + 1)


def riga_totale_scheda(ws, col0: int):
    return ws.cell(row=ex._Esportatore.S_R_TOT + 1, column=col0 + 1)


def documento(**kw) -> Document:
    base = {
        "id": "x1",
        "source_file": "prova.pdf",
        "status": "completato",
        "header": Header(operatore="ROSSI MARIO", mese=2, anno=2026),
    }
    base.update(kw)
    return Document(**base)


# --- Struttura ------------------------------------------------------------------


def test_fogli_nell_ordine_previsto(wb_formule):
    nomi = wb_formule.sheetnames
    assert nomi[:2] == ["Riepilogo", "Dettaglio giornaliero"]
    assert nomi[-3:] == ["Anomalie", "Totali per operatore", "Legenda e note"]
    assert schede(wb_formule) == [
        "DE LUCA G. 03-2026",
        "ESPOSITO A. 02-2026",
        "ROSSI M. 02-2026",
        "ROSSI M. 03-2026",
    ]


def test_nomi_schede_validi_e_univoci():
    docs = []
    for i, op in enumerate(
        [
            "ROSSI MARIO",
            "ROSSI MARIO",
            "rossi mario",
            "D'ANGELO MARIA",
            "A" * 60,
            "BIANCHI [LUCA]: *?/\\",
            None,
            "History",
            "DELLA VALLE SANTORO MARIA GRAZIA ANTONIETTA",
            "ROSSI MARIO LUIGI",
        ]
    ):
        docs.append(documento(id=f"d{i}", header=Header(operatore=op, mese=2, anno=2026), source_page=i + 1))
    prepared = [ex._prepara(d, i) for i, d in enumerate(docs, start=1)]
    usati = {s.casefold() for s in ex._FOGLI_FISSI}
    nomi = [ex._nome_scheda(d, usati) for d in prepared]
    assert nomi[:4] == ["ROSSI M. 02-2026", "ROSSI M. 02-2026 (2)", "rossi m. 02-2026 (3)", "D'ANGELO M. 02-2026"]
    assert nomi[6] == "Foglio 7 02-2026"
    assert nomi[8:] == ["DELLA VALLE S.M.G.A. 02-2026", "ROSSI MARIO L. 02-2026"]
    assert len({n.casefold() for n in nomi}) == len(nomi)
    for n in nomi:
        assert 1 <= len(n) <= 31
        assert not RE_NOME_VIETATO.search(n)
        assert not n.startswith("'") and not n.endswith("'")
        assert n.casefold() != "history"


def test_apostrofo_nel_nome_della_scheda(tmp_path):
    rows = empty_rows()
    rows[1] = DayRow(
        giorno=2,
        prog_entrata="8:00",
        prog_uscita="11:00",
        eff_entrata="8:00",
        eff_uscita="11:00",
        ore_dichiarate=3,
        firma=True,
    )
    doc = documento(
        header=Header(operatore="D'ANGELO MARIA", mese=2, anno=2026, totale_mensile_dichiarato=3), rows=rows
    )
    path = export_workbook([doc], tmp_path / "apostrofo.xlsx")
    wb = openpyxl.load_workbook(path)
    assert "D'ANGELO M. 02-2026" in wb.sheetnames
    formula = wb["Riepilogo"]["J12"].value
    assert formula == "='D''ANGELO M. 02-2026'!$O$45"
    assert wb["Riepilogo"]["B12"].hyperlink.location == "'D''ANGELO M. 02-2026'!A1"


# --- Tipi dei valori ---------------------------------------------------------------


def test_orari_date_e_ore_sono_valori_numerici(wb_formule):
    ws = wb_formule["ROSSI M. 02-2026"]
    E = ex._Esportatore
    entrata = cella_scheda(ws, 2, E.S_PE)
    assert isinstance(entrata.value, (datetime, time, float))
    assert entrata.number_format == "hh:mm"
    assert entrata.data_type == "d" or entrata.data_type == "n"
    data = cella_scheda(ws, 2, E.S_DATA)
    assert isinstance(data.value, datetime) and data.value.date().isoformat() == "2026-02-02"
    assert data.number_format == "dd/mm/yyyy"
    ore = cella_scheda(ws, 4, E.S_OD)
    assert ore.value == 1.5 and ore.number_format == ex.FMT_ORE
    assert cella_scheda(ws, 2, E.S_WD).value == "lun"
    mese = ws.cell(row=5, column=12)
    assert isinstance(mese.value, datetime) and mese.number_format == ex.FMT_MESE
    assert ws.cell(row=9, column=4).value == 1  # Lotto come numero


def test_nessun_numero_memorizzato_come_testo(wb_formule):
    sospetti = []
    for ws in wb_formule.worksheets:
        intestazioni = {
            c.column: str(c.value) for row in ws.iter_rows(max_row=13) for c in row if isinstance(c.value, str)
        }
        for row in ws.iter_rows():
            for c in row:
                if (
                    isinstance(c.value, str)
                    and not c.value.startswith("=")
                    and RE_NUMERO_COME_TESTO.match(c.value)
                ):
                    # le note libere sono testo per natura (es. "104" = permesso L. 104)
                    if "Note" in intestazioni.get(c.column, ""):
                        continue
                    sospetti.append(f"{ws.title}!{c.coordinate}={c.value!r}")
    assert not sospetti, sospetti


def test_orari_nel_dettaglio(wb_formule):
    ws = wb_formule["Dettaglio giornaliero"]
    intest = {c.value: c.column for c in ws[5]}
    for nome in ("Entrata progr.", "Uscita progr.", "Entrata eff.", "Uscita eff."):
        c = ws.cell(row=6, column=intest[nome])
        assert c.number_format == "hh:mm" and not isinstance(c.value, str)
    assert ws.cell(row=6, column=intest["Data"]).number_format == "dd/mm/yyyy"
    assert isinstance(ws.cell(row=6, column=intest["Data"]).value, datetime)
    assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A5:")
    assert ws.freeze_panes == "D6"


# --- Formule e totali ----------------------------------------------------------------


def test_totali_delle_schede_coincidono_con_la_validazione(esempio, wb_formule, wb_valori):
    _, docs = esempio
    attesi = validati(docs)
    per_scheda = {
        "ROSSI M. 02-2026": attesi["a1b2c3d4e5f6"],
        "ESPOSITO A. 02-2026": attesi["b2c3d4e5f6a1"],
        "ROSSI M. 03-2026": attesi["c3d4e5f6a1b2"],
        "DE LUCA G. 03-2026": attesi["d4e5f6a1b2c3"],
    }
    E = ex._Esportatore
    for nome, doc in per_scheda.items():
        wf, wv = wb_formule[nome], wb_valori[nome]
        t = doc.totals
        for col, atteso in (
            (E.S_OD, t.ore_dichiarate),
            (E.S_OP, t.ore_programmate),
            (E.S_OC, t.ore_calcolate),
            (E.S_OR, t.ore_riconosciute),
            (E.S_AA, t.assenze_alunno),
            (E.S_AO, t.assenze_operatore),
        ):
            assert riga_totale_scheda(wf, col).value.startswith(("=SUM(", "=COUNTIF(")), (nome, col)
            assert riga_totale_scheda(wv, col).value == pytest.approx(atteso), (nome, col)
        # ore calcolate per riga: formula (uscita - entrata) * 24 con controlli sulle celle vuote
        f = cella_scheda(wf, 2, E.S_OC).value
        assert f.startswith("=IF(AND(ISNUMBER(") and "*24" in f
        giorni = trova_riga(wf, 1, "Giorni lavorati (con ore riconosciute)")
        assert wv.cell(row=giorni, column=8).value == t.giorni_lavorati


def test_confronto_con_il_totale_dichiarato(wb_formule, wb_valori):
    ws_f, ws_v = wb_formule["ROSSI M. 03-2026"], wb_valori["ROSSI M. 03-2026"]
    r = trova_riga(ws_f, 1, "Differenza (somma dei giorni − totale dichiarato)")
    assert ws_f.cell(row=r, column=8).value.startswith("=IF(ISNUMBER(")
    assert ws_v.cell(row=r, column=8).value == pytest.approx(0.5)
    assert ws_v.cell(row=r + 1, column=8).value == "Non corrisponde"
    assert wb_valori["ROSSI M. 02-2026"].cell(row=r + 1, column=8).value == "Corrisponde"


def test_riepilogo_tabella_formule_e_totali(esempio, wb_formule, wb_valori):
    _, docs = esempio
    attesi = validati(docs)
    ws, wv = wb_formule["Riepilogo"], wb_valori["Riepilogo"]
    assert "Riepilogo_fogli" in ws.tables
    tabella = ws.tables["Riepilogo_fogli"]
    assert tabella.ref == "A11:W15"
    intest = {c.value: c.column for c in ws[11]}
    # una riga per documento, collegata alla scheda
    for r in range(12, 16):
        cella = ws.cell(row=r, column=intest["Operatore"])
        assert cella.hyperlink is not None and cella.hyperlink.location.endswith("!A1")
        assert ws.cell(row=r, column=intest["Ore riconosciute"]).value.startswith("='")
    # riga dei totali con SUBTOTAL (rispetta i filtri)
    tot = ws.cell(row=16, column=intest["Ore riconosciute"])
    assert tot.value == "=SUBTOTAL(109,$J$12:$J$15)"
    atteso = sum(d.totals.ore_riconosciute for d in attesi.values())
    assert wv.cell(row=16, column=intest["Ore riconosciute"]).value == pytest.approx(atteso)
    assert wv.cell(row=16, column=intest["Ore programmate"]).value == pytest.approx(
        sum(d.totals.ore_programmate for d in attesi.values())
    )
    assert wv.cell(row=16, column=intest["Errori"]).value == sum(d.totals.n_errori for d in attesi.values())
    # stati
    stati = [wv.cell(row=r, column=intest["Stato"]).value for r in range(12, 16)]
    assert sorted(stati) == ["Da verificare", "Errori", "Errori", "OK"]
    # indicatori in alto: formule collegate alla riga dei totali
    assert ws["A6"].value == "=SUBTOTAL(103,$B$12:$B$15)" and wv["A6"].value == 4
    assert wv.cell(row=12, column=intest["Verificato"]).value in ("Sì", "No")
    assert ws.freeze_panes == "C12"


def test_differenza_totale_nel_riepilogo(wb_valori):
    ws = wb_valori["Riepilogo"]
    intest = {c.value: c.column for c in ws[11]}
    r = next(
        r for r in range(12, 16) if ws.cell(row=r, column=intest["File di origine"]).value == "rossi_marzo.pdf"
    )
    assert ws.cell(row=r, column=intest["Totale mensile dichiarato"]).value == 64
    assert ws.cell(row=r, column=intest["Differenza totale"]).value == pytest.approx(0.5)


def test_totali_per_operatore_con_somma_piu_se(esempio, wb_formule, wb_valori):
    _, docs = esempio
    attesi = validati(docs)
    ws, wv = wb_formule["Totali per operatore"], wb_valori["Totali per operatore"]
    r = trova_riga(ws, 1, "ROSSI MARIO")
    intest = {c.value: c.column for c in ws[trova_riga(ws, 1, "Operatore")]}
    formula = ws.cell(row=r, column=intest["Totale ore riconosciute"]).value
    assert formula.startswith("=SUMIFS('Dettaglio giornaliero'!")
    rossi = [d for d in attesi.values() if d.header.operatore == "ROSSI MARIO"]
    assert wv.cell(row=r, column=intest["Totale ore riconosciute"]).value == pytest.approx(
        sum(d.totals.ore_riconosciute for d in rossi)
    )
    assert wv.cell(row=r, column=intest["Ore febbraio 2026"]).value == pytest.approx(52.5)
    assert wv.cell(row=r, column=intest["Ore marzo 2026"]).value == pytest.approx(64.5)
    assert wv.cell(row=r, column=intest["Fogli"]).value == 2
    assert wv.cell(row=r, column=intest["Giorni lavorati"]).value == sum(d.totals.giorni_lavorati for d in rossi)


def test_dettaglio_una_riga_per_giorno_con_dati(esempio, wb_valori):
    _, docs = esempio
    attesi = validati(docs)
    ws = wb_valori["Dettaglio giornaliero"]
    righe = [r for r in range(6, ws.max_row + 1) if isinstance(ws.cell(row=r, column=1).value, int)]
    attese = sum(1 for d in attesi.values() for r in d.rows if r.has_content())
    assert len(righe) == attese
    intest = {c.value: c.column for c in ws[5]}
    somma = sum(ws.cell(row=r, column=intest["Ore riconosciute"]).value or 0 for r in righe)
    assert somma == pytest.approx(sum(d.totals.ore_riconosciute for d in attesi.values()))


def test_opzione_giorni_vuoti(tmp_path):
    docs = documenti_esempio()
    path = export_workbook(docs, tmp_path / "vuoti.xlsx", ExportOptions(giorni_vuoti=True))
    ws = openpyxl.load_workbook(path)["Dettaglio giornaliero"]
    righe = [r for r in range(6, ws.max_row + 1) if isinstance(ws.cell(row=r, column=1).value, int)]
    assert len(righe) == 28 + 28 + 31 + 31


# --- Evidenziazioni: illeggibile, incerto, corretto -------------------------------------


def _colore(c) -> str:
    return (c.fill.fgColor.rgb or "")[-6:].upper()


def test_cella_illeggibile_evidenziata_e_commentata(wb_formule):
    ws = wb_formule["ROSSI M. 02-2026"]
    c = cella_scheda(ws, 12, ex._Esportatore.S_EU)
    assert c.value == "ILLEGGIBILE"
    assert _colore(c) == ex.RED_FILL[1:].upper()
    assert c.font.b and (c.font.color.rgb or "")[-6:].upper() == ex.RED_DARK[1:].upper()
    assert c.comment is not None and "illeggibile" in c.comment.text.lower()
    # anche nel Dettaglio
    wd = wb_formule["Dettaglio giornaliero"]
    trovate = [x for row in wd.iter_rows() for x in row if x.value == "ILLEGGIBILE"]
    assert len(trovate) == 1 and trovate[0].comment is not None


def test_cella_incerta_evidenziata_e_commentata(wb_formule):
    ws = wb_formule["ROSSI M. 02-2026"]
    c = cella_scheda(ws, 18, ex._Esportatore.S_OD)
    assert c.value == 3
    assert _colore(c) == ex.AMBER_FILL[1:].upper()
    assert c.comment is not None and c.comment.text.startswith("Lettura incerta: verificare sull'originale.")
    assert "55%" in c.comment.text
    alunno = ws.cell(row=8, column=4)
    assert alunno.value == "BIANCHI LUCA" and _colore(alunno) == ex.AMBER_FILL[1:].upper()


def test_cella_corretta_a_mano_in_blu_con_valore_originale(wb_formule):
    ws = wb_formule["ESPOSITO A. 02-2026"]
    c = cella_scheda(ws, 5, ex._Esportatore.S_EE)
    assert (c.font.color.rgb or "")[-6:].upper() == ex.ACCENT_DARK[1:].upper()
    assert c.comment is not None
    assert c.comment.text == "Corretto manualmente – valore OCR originale: 11:30"


def test_weekend_e_giorni_inesistenti(wb_formule):
    ws = wb_formule["ROSSI M. 02-2026"]
    E = ex._Esportatore
    assert _colore(cella_scheda(ws, 1, E.S_G)) == ex.WEEKEND[1:].upper()  # domenica 1/2/2026
    assert _colore(cella_scheda(ws, 30, E.S_G)) == ex.INESISTENTE[1:].upper()
    assert cella_scheda(ws, 30, E.S_DATA).value is None
    assert cella_scheda(ws, 30, E.S_ES).value == "Giorno inesistente nel mese"


# --- Collegamenti, anomalie, legenda --------------------------------------------------


def test_collegamenti_ipertestuali(wb_formule):
    for nome in schede(wb_formule):
        ws = wb_formule[nome]
        links = [c.hyperlink.location for row in ws.iter_rows() for c in row if c.hyperlink]
        assert "'Riepilogo'!A1" in links
    wa = wb_formule["Anomalie"]
    links = [c.hyperlink.location for row in wa.iter_rows() for c in row if c.hyperlink]
    assert any(loc == "'ROSSI M. 02-2026'!A25" for loc in links)  # giorno 12 -> riga 25


def test_foglio_anomalie_completo(esempio, wb_valori):
    _, docs = esempio
    attesi = validati(docs)
    ws = wb_valori["Anomalie"]
    intest = {c.value: c.column for c in ws[5]}
    righe = [r for r in range(6, ws.max_row + 1) if isinstance(ws.cell(row=r, column=1).value, int)]
    assert len(righe) == sum(len(d.anomalies) for d in attesi.values())
    codici = {ws.cell(row=r, column=intest["Codice"]).value for r in righe}
    assert {
        "E01_ORE_NON_COERENTI",
        "E02_TOTALE_MENSILE_DIVERSO",
        "E05_CAMPO_ILLEGGIBILE",
        "W04_CAMPO_INCERTO",
        "I02_ASSENZA_OPERATORE",
    } <= codici
    gravita = {ws.cell(row=r, column=intest["Gravità"]).value for r in righe}
    assert gravita == {"Errore", "Attenzione", "Info"}


def test_legenda_elenca_documenti_non_inclusi(wb_valori):
    ws = wb_valori["Legenda e note"]
    testi = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert "lettera_trasmissione.pdf" in testi
    assert any(t.startswith("Elaborazione non riuscita: Connessione") for t in testi)
    assert "Documenti non inclusi (2)" in testi
    assert any("claude-opus-5-5" in t for t in testi)


def test_impostazioni_di_stampa(wb_formule):
    for ws in wb_formule.worksheets:
        assert ws.page_setup.orientation == "landscape"
        assert ws.page_setup.paperSize == 9
        assert ws.sheet_properties.pageSetUpPr.fitToPage
        assert ws.page_setup.fitToWidth in (None, 1)  # 1 è il valore predefinito (attributo omesso)
        assert ws.page_setup.fitToHeight == 0
        assert ws.sheet_view.showGridLines is False
    assert wb_formule["ROSSI M. 02-2026"].print_title_rows == "$12:$13"


# --- Casi limite ---------------------------------------------------------------------


def test_nulla_da_esportare():
    with pytest.raises(ValueError, match="Nessun documento"):
        export_workbook([], Path("non_creato.xlsx"))
    in_coda = documento(status="in_coda")
    scartato = documento(id="x2", status="completato", is_foglio_firma=False)
    with pytest.raises(ValueError, match="Nessun foglio firma completato"):
        export_workbook([in_coda, scartato], Path("non_creato.xlsx"))
    assert not Path("non_creato.xlsx").exists()


def test_i_documenti_ricevuti_non_vengono_modificati(tmp_path):
    docs = documenti_esempio()
    prima = [d.model_dump() for d in docs]
    export_workbook(docs, tmp_path / "x.xlsx")
    assert [d.model_dump() for d in docs] == prima


def test_estensione_aggiunta_e_nessun_file_temporaneo(tmp_path):
    path = export_workbook(documenti_esempio(), tmp_path / "rendiconto")
    assert path.name == "rendiconto.xlsx" and path.exists()
    assert sorted(p.name for p in tmp_path.iterdir()) == ["rendiconto.xlsx"]


def test_documento_senza_mese_e_con_giorno_inesistente(tmp_path):
    rows = empty_rows()
    rows[0] = DayRow(giorno=1, eff_entrata="9", eff_uscita="12", ore_dichiarate=3, firma=True)
    senza_mese = documento(id="s1", header=Header(operatore="VERDI PAOLO"), rows=rows)
    rows2 = empty_rows()
    rows2[29] = DayRow(giorno=30, eff_entrata="08:00", eff_uscita="11:00", ore_dichiarate=3, firma=True)
    rows2[1] = DayRow(giorno=2, eff_entrata="08:00", eff_uscita="11:00", ore_dichiarate=3, firma=True)
    feb = documento(id="s2", header=Header(operatore="NERI ANNA", mese=2, anno=2026), rows=rows2)
    path = export_workbook([senza_mese, feb], tmp_path / "limiti.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    assert "VERDI P." in wb.sheetnames
    ws = wb["NERI A. 02-2026"]
    E = ex._Esportatore
    assert riga_totale_scheda(ws, E.S_OR).value == 3  # il 30 febbraio non è conteggiato
    assert cella_scheda(ws, 30, E.S_ES).value.startswith("Errore – giorno inesistente")
    assert wb["VERDI P."].cell(row=E.S_R_D1 + 1, column=E.S_OC + 1).value == 3


def test_testo_che_inizia_con_uguale_resta_testo(tmp_path):
    rows = empty_rows()
    rows[1] = DayRow(giorno=2, note='=HYPERLINK("http://example.com","x")')
    doc = documento(rows=rows)
    path = export_workbook([doc], tmp_path / "formula.xlsx")
    ws = openpyxl.load_workbook(path)["ROSSI M. 02-2026"]
    c = cella_scheda(ws, 2, ex._Esportatore.S_NO)
    assert c.data_type == "s" and c.value.startswith("=HYPERLINK")


def test_senza_schede_per_documento(tmp_path):
    path = export_workbook(
        documenti_esempio(), tmp_path / "compatto.xlsx", ExportOptions(fogli_per_documento=False)
    )
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == FISSI[:2] + FISSI[2:]
    ws = wb["Riepilogo"]
    assert ws["J12"].value.startswith("=SUMIFS('Dettaglio giornaliero'!")
    assert ws["B12"].hyperlink is None


def test_titolo_personalizzato(tmp_path):
    path = export_workbook(
        documenti_esempio(), tmp_path / "t.xlsx", ExportOptions(titolo="Rendiconto marzo – Lotto 1")
    )
    wb = openpyxl.load_workbook(path)
    assert wb["Riepilogo"]["A1"].value == "Rendiconto marzo – Lotto 1"
    assert wb.properties.title.startswith("Rendiconto marzo – Lotto 1")


def test_nome_file_predefinito(monkeypatch):
    monkeypatch.setattr(ex, "_adesso", lambda: datetime(2026, 10, 9, 15, 30))  # noqa: DTZ001 (ora locale)
    docs = documenti_esempio()
    assert default_filename(docs) == "Rendicontazione_2026-02_2026-03_20261009-1530.xlsx"
    feb = [d for d in docs if d.header.mese == 2]
    assert default_filename(feb) == "Rendicontazione_2026-02_20261009-1530.xlsx"
    assert default_filename([]) == "Rendicontazione_20261009-1530.xlsx"


def test_partizione_dei_riquadri():
    gruppi = ex._partizione([5, 24, 24, 22, 13, 8, 10, 10, 10, 10, 11, 10], 4)
    assert len(gruppi) == 4 and gruppi[0][0] == 0 and gruppi[-1][1] == 11
    assert all(b + 1 == c for (_, b), (c, _) in itertools.pairwise(gruppi))


# --- Ricalcolo con LibreOffice -----------------------------------------------------------

_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema"
 xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def _ricalcola(src: Path, cartella: Path) -> Path:
    """Apre il file con LibreOffice forzando il ricalcolo di tutte le formule e lo salva in ``cartella``."""
    profilo = cartella / "profilo"
    (profilo / "user").mkdir(parents=True, exist_ok=True)
    (profilo / "user" / "registrymodifications.xcu").write_text(_XCU, encoding="utf-8")
    out = cartella / "ricalcolato"
    subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation={profilo.as_uri()}",
            "--headless",
            "--calc",
            "--convert-to",
            "xlsx:Calc MS Excel 2007 XML",
            "--outdir",
            str(out),
            str(src),
        ],
        check=True,
        capture_output=True,
        timeout=300,
    )
    return out / src.name


def _confronta_con_libreoffice(path: Path, tmp_path: Path) -> openpyxl.Workbook:
    # verifica preliminare: il profilo forza davvero il ricalcolo
    canarino = tmp_path / "canarino.xlsx"
    wb = xlsxwriter.Workbook(str(canarino))
    wb.add_worksheet().write_formula("A1", "=1+1", None, 99)
    wb.close()
    assert openpyxl.load_workbook(_ricalcola(canarino, tmp_path), data_only=True).active["A1"].value == 2

    ricalcolato = openpyxl.load_workbook(_ricalcola(path, tmp_path), data_only=True)
    formule = openpyxl.load_workbook(path, data_only=False)
    memorizzati = openpyxl.load_workbook(path, data_only=True)
    n, diverse, errori = 0, [], []
    for ws in formule.worksheets:
        for row in ws.iter_rows():
            for c in row:
                lo = ricalcolato[ws.title][c.coordinate].value
                if isinstance(lo, str) and lo.startswith("#"):
                    errori.append(f"{ws.title}!{c.coordinate}: {lo}")
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    continue
                n += 1
                mio = memorizzati[ws.title][c.coordinate].value
                if mio in ("", None) and lo in ("", None):
                    continue
                if isinstance(mio, (int, float)) and isinstance(lo, (int, float)):
                    if abs(mio - lo) > 1e-9:
                        diverse.append(f"{ws.title}!{c.coordinate}: {mio} != {lo}")
                elif mio != lo:
                    diverse.append(f"{ws.title}!{c.coordinate}: {mio!r} != {lo!r}")
    assert n > 100
    assert not errori, errori[:20]
    assert not diverse, diverse[:20]
    return ricalcolato


requires_soffice = pytest.mark.skipif(
    shutil.which("soffice") is None, reason="LibreOffice (soffice) non installato"
)


@requires_soffice
def test_libreoffice_ricalcola_gli_stessi_valori(esempio, tmp_path):
    path, docs = esempio
    ricalcolato = _confronta_con_libreoffice(path, tmp_path)
    attesi = validati(docs)
    ws = ricalcolato["Riepilogo"]
    assert ws["J16"].value == pytest.approx(sum(d.totals.ore_riconosciute for d in attesi.values()))
    assert (
        ricalcolato["ROSSI M. 03-2026"]
        .cell(row=ex._Esportatore.S_R_TOT + 1, column=ex._Esportatore.S_OR + 1)
        .value
        == 64.5
    )


@requires_soffice
def test_libreoffice_senza_schede(tmp_path):
    path = export_workbook(
        documenti_esempio(), tmp_path / "compatto.xlsx", ExportOptions(fogli_per_documento=False)
    )
    _confronta_con_libreoffice(path, tmp_path)


# --- Foglio firma reale (facoltativo) -------------------------------------------------------


def test_verita_del_foglio_reale(tmp_path):
    verita = os.environ.get("SIRIO_SAMPLE_TRUTH")
    if not verita or not Path(verita).exists():
        pytest.skip("SIRIO_SAMPLE_TRUTH non impostata")
    dati = json.loads(Path(verita).read_text(encoding="utf-8"))
    doc = Document(
        id="campione",
        source_file="campione.pdf",
        status="completato",
        header=Header(**dati["header"]),
        rows=[DayRow(**r) for r in dati["rows"]],
    )
    path = export_workbook([doc], tmp_path / "campione.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    scheda = schede(wb)[0]
    assert re.fullmatch(r".+ [A-Z]\. 02-2026", scheda)
    ws = wb[scheda]
    E = ex._Esportatore
    totale = float(dati["header"]["totale_mensile_dichiarato"])
    assert riga_totale_scheda(ws, E.S_OD).value == pytest.approx(totale)
    assert riga_totale_scheda(ws, E.S_OR).value == pytest.approx(totale)
    assert wb["Riepilogo"]["J13"].value == pytest.approx(totale)


def _documenti_casi_limite() -> list[Document]:
    """Documento senza mese, febbraio con dati il 30, totale illeggibile, orario non valido, assenze."""
    rows = empty_rows()
    rows[0] = DayRow(giorno=1, eff_entrata="9", eff_uscita="12", ore_dichiarate=3, firma=True)
    rows[1] = DayRow(giorno=2, eff_entrata="8:70", eff_uscita="11:00", ore_dichiarate=3, firma=True)
    senza_mese = documento(id="s1", header=Header(operatore="VERDI PAOLO", ore_pei=10), rows=rows)
    rows2 = empty_rows()
    rows2[1] = DayRow(
        giorno=2,
        prog_entrata="08:00",
        prog_uscita="11:00",
        eff_entrata="08:00",
        eff_uscita="11:00",
        ore_dichiarate=3,
        firma=True,
    )
    rows2[2] = DayRow(giorno=3, prog_entrata="08:00", prog_uscita="11:00", assenza_operatore=True, note="104")
    rows2[3] = DayRow(
        giorno=4,
        prog_entrata="08:00",
        prog_uscita="11:00",
        eff_entrata="08:00",
        eff_uscita="11:00",
        ore_dichiarate=1.5,
        assenza_alunno=True,
        firma=True,
    )
    rows2[4] = DayRow(
        giorno=5,
        prog_entrata="08:00",
        eff_entrata="08:00",
        firma=True,
        illeggibili=["eff_uscita", "ore_dichiarate"],
    )
    rows2[29] = DayRow(giorno=30, eff_entrata="08:00", eff_uscita="11:00", ore_dichiarate=3, firma=True)
    feb = documento(
        id="s2",
        header=Header(
            operatore="NERI ANNA", mese=2, anno=2026, ore_pei=6, illeggibili=["totale_mensile_dichiarato"]
        ),
        rows=rows2,
    )
    return [senza_mese, feb]


@requires_soffice
def test_libreoffice_casi_limite(tmp_path):
    path = export_workbook(_documenti_casi_limite(), tmp_path / "limiti.xlsx", ExportOptions(giorni_vuoti=True))
    ricalcolato = _confronta_con_libreoffice(path, tmp_path)
    ws = ricalcolato["NERI A. 02-2026"]
    r = trova_riga(ws, 1, "Esito del confronto")
    assert ws.cell(row=r, column=8).value == "Totale non leggibile"
    assert ws.cell(row=r - 3, column=8).value == "ILLEGGIBILE"
    tot = ricalcolato["Totali per operatore"]
    r = trova_riga(tot, 1, "VERDI PAOLO")
    intest = {c.value: c.column for c in tot[trova_riga(tot, 1, "Operatore")]}
    assert tot.cell(row=r, column=intest["Ore mese non indicato"]).value == pytest.approx(6)
    assert tot.cell(row=r, column=intest["Ore febbraio 2026"]).value == 0
